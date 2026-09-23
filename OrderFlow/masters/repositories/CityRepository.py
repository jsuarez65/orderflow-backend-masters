from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError

from configuration.DatabaseConfiguration import sessionLocal
from configuration.LogConfiguration import LogConfiguration
from model.entities.CityEntity import CityEntity

class CityRepository:
    def __init__(self):
        self.log = LogConfiguration.getLogger()
        self.session = sessionLocal()

    def importPostalCodes(self, file):
        workbook = load_workbook(file)
        sheet = workbook.active

        inserted = 0
        ignored_duplicates = 0
        skipped = 0

        try:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row or row[0] is None:
                    skipped += 1
                    continue

                try:
                    raw_cp = row[0]
                    if isinstance(raw_cp, float):
                        postal_code = str(int(raw_cp))
                    else:
                        postal_code = str(raw_cp).strip()

                    if not postal_code:
                        skipped += 1
                        continue

                    city_name = str(row[1]).strip().title() if row[1] and str(row[1]).strip() else ""

                    city_entity = CityEntity(
                        codigo_postal=postal_code,
                        nombre_localidad=city_name
                    )

                    self.session.add(city_entity)
                    self.session.commit()
                    inserted += 1

                except IntegrityError:
                    self.session.rollback()
                    ignored_duplicates += 1
                except Exception as e:
                    self.session.rollback()
                    skipped += 1
                    self.log.error(f"Error in row {row}: {e}")
                    continue

            return {
                "message": "Import finished successfully",
                "inserted_new": inserted,
                "ignored_duplicates": ignored_duplicates,
                "skipped_rows": skipped,
                "total_processed": inserted + ignored_duplicates + skipped
            }
        finally:
            self.session.close()
    
    def getPostalCodes(self, postal_code: str = None, city_name: str = None):
        query = self.session.query(CityEntity)

        if postal_code:
            query = query.filter(CityEntity.codigo_postal == postal_code.strip())
        
        if city_name:
            query = query.filter(CityEntity.nombre_localidad.ilike(f"%{city_name.strip()}%"))

        cities = query.all()

        return [
            {
                "id": city.id,
                "postal_code": city.codigo_postal,
                "city_name": city.nombre_localidad
            }
            for city in cities
        ]