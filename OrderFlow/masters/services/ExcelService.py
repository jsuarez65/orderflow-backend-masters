from openpyxl import load_workbook

class ExcelService:
    def __init__(self, log):
        self.log = log

    def readPostalCodes(self, file):
        workbook = load_workbook(file)
        sheet = workbook.active

        postalCodes = []

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or row[0] is None:
                continue

            rawCp = row[0]
            postalCode = str(int(rawCp)) if isinstance(rawCp, float) else str(rawCp).strip()

            if not postalCode:
                continue

            cityName = str(row[1]).strip().title() if row[1] and str(row[1]).strip() else ""

            postalCodes.append((postalCode, cityName))

        return postalCodes